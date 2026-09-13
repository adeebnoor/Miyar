import copy,base64,io,json
import pytest
from sqlalchemy import select,text
from sqlalchemy.exc import DBAPIError
from server.models import Position,User,OutboxEvent,Organization
from server.domain import DEFAULT_FRAMEWORK,grade,validate_framework
from server.security import issue_token
from .conftest import submit,decide,activate

def test_auth_tenant_department_and_claim_spoofing(env):
    app,c,auth,create=env;p=create();url='/api/v1/positions/'+p['id']
    assert c.get(url).status_code==401
    assert c.get(url,headers=auth('admin','org-b')).status_code==404
    assert c.get('/api/v1/positions',headers=auth('line_manager','org-b')).json()['total']==0
    with app.state.sessions() as db:
        manager=db.get(User,'org-a-line_manager');manager.department_id='org-a-two';db.commit()
    assert c.get(url,headers=auth()).status_code==404
    assert c.patch(url,headers=auth('finance'),json={'revision':1,'content':p['content'],'reason':'Unauthorized change'}).status_code==403
    assert c.post('/api/v1/positions',headers=auth('admin'),json={'departmentId':'org-b-one','content':p['content'],'reason':'Cross tenant'}).status_code==403
    assert c.post('/api/v1/positions',headers=auth('admin'),json={'departmentId':'org-a-one','content':{**p['content'],'state':'active'},'reason':'State injection'}).status_code==422

def test_all_stages_budget_and_duplicate_approval(env):
    app,c,auth,create=env;p=create();submit(c,auth,p)
    assert decide(c,auth,p,'chro',{}).status_code==403
    assert decide(c,auth,p,'admin',{}).status_code==403
    assert decide(c,auth,p,'od_specialist',{}).status_code==422
    assert decide(c,auth,p,'od_specialist',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01'}).status_code==200
    assert decide(c,auth,p,'od_specialist',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01'}).status_code==403
    assert decide(c,auth,p,'total_rewards',{'payFrameworkReviewed':True}).status_code==422
    assert c.post('/api/v1/positions/'+p['id']+'/evaluation',headers=auth('total_rewards'),json={'revision':1,'answers':{'knowledge':'2','complexity':'2','impact':'2'},'evidence':dict.fromkeys(['knowledge','complexity','impact'],'Scoped role evidence')}).status_code==200
    assert decide(c,auth,p,'total_rewards',{'payFrameworkReviewed':True}).status_code==200
    for count,budget in [(True,240000),(1,239999),(0,240000),(1.5,240000)]:
        assert decide(c,auth,p,'finance',{'vacancyConfirmed':True,'budgetConfirmed':True,'approvedAnnualBudget':budget,'approvedHeadcount':count}).status_code==422
    assert decide(c,auth,p,'finance',{'vacancyConfirmed':True,'budgetConfirmed':True,'approvedAnnualBudget':240000,'approvedHeadcount':1}).status_code==200
    assert decide(c,auth,p,'chro',{}).json()['activeRevision']==1
    with app.state.sessions() as db:assert len(list(db.scalars(select(OutboxEvent))))==1

def test_requester_cannot_self_approve_even_after_role_change(env):
    app,c,auth,create=env;p=create();submit(c,auth,p)
    with app.state.sessions() as db:u=db.get(User,'org-a-line_manager');u.role='od_specialist';db.commit()
    assert decide(c,auth,p,'line_manager',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01'}).status_code==403

def test_active_revision_survives_amendment_and_restore(env):
    app,c,auth,create=env;p=activate(c,auth,create());url='/api/v1/positions/'+p['id'];new={**p['content'],'title':'Revised title','annualCost':999999}
    r=c.patch(url,headers=auth(),json={'revision':1,'content':new,'reason':'New scope proposal'});assert r.status_code==200,r.text;assert r.json()['activeRevision']==1
    assert c.get('/api/v1/analytics',headers=auth('chro')).json()['annualPositionCost']==240000
    assert c.get(url+'/export/json',headers=auth()).json()['content']['title']=='مهندس برمجيات'
    assert c.patch(url,headers=auth(),json={'revision':1,'content':new,'reason':'Stale tab write'}).status_code==409
    r=c.post(url+'/restore',headers=auth(),json={'revision':2,'restoreRevision':1,'reason':'Restore prior scope'});assert r.json()['revision']==3
    assert len(c.get(url+'/versions',headers=auth()).json())==3

def test_edit_in_review_and_return_require_new_revision(env):
    app,c,auth,create=env;p=create();submit(c,auth,p);url='/api/v1/positions/'+p['id']
    assert c.patch(url,headers=auth(),json={'revision':1,'content':p['content'],'reason':'Change under review'}).status_code==409
    assert decide(c,auth,p,'od_specialist',{},'return').status_code==200
    assert c.post(url+'/submit',headers=auth(),json={'revision':1,'reason':'Skip version change'}).status_code==409
    r=c.patch(url,headers=auth(),json={'revision':1,'content':p['content'],'reason':'Answer reviewer concern'});assert r.json()['revision']==2

def test_audit_and_versions_reject_sql_update_delete(env):
    app,c,auth,create=env;p=create();assert c.get('/api/v1/audit/verify',headers=auth('admin')).json()['valid'] is True
    for table in ['audit_events','position_versions']:
        for query in [f"UPDATE {table} SET id=id",f"DELETE FROM {table}"]:
            with app.state.sessions() as db:
                with pytest.raises(DBAPIError):db.execute(text(query));db.commit()
    assert c.get('/api/v1/audit/verify',headers=auth('admin')).json()['valid'] is True

def test_illustrative_and_proprietary_grade_gates(env):
    app,c,auth,create=env
    f=copy.deepcopy(DEFAULT_FRAMEWORK);f.update(method='korn-ferry-licensed',licenseConfirmed=True,licenseReference='test')
    assert c.post('/api/v1/settings/framework',headers=auth('admin'),json={'framework':f,'reason':'Validate method provenance'}).status_code==422
    assert grade(DEFAULT_FRAMEWORK,dict.fromkeys(['knowledge','complexity','impact'],'2'),dict.fromkeys(['knowledge','complexity','impact'],'Evidence'))['points']==500
    f=copy.deepcopy(DEFAULT_FRAMEWORK);f['bands'][1]['min']=300
    with pytest.raises(ValueError):validate_framework(f)

def test_provisional_cannot_be_a_fake_national_code(env):
    app,c,auth,create=env;p=create(occupationCode='',provisional=True,mappingJustification='Hybrid position pending taxonomy review');p=activate(c,auth,p)
    result=c.get('/api/v1/positions/'+p['id']+'/export/json',headers=auth()).json();assert not result['content']['occupationCode'];assert result['internalCode'].startswith('MJR-')

def test_export_is_scoped_editable_and_signed(env,monkeypatch):
    app,c,auth,create=env;monkeypatch.setenv('MIYAR_SIGNING_KEY',base64.b64encode(b'0123456789abcdef0123456789abcdef').decode());p=activate(c,auth,create());url='/api/v1/positions/'+p['id']+'/export/'
    assert c.get(url+'pdf',headers=auth('admin','org-b')).status_code==404
    r=c.get(url+'receipt',headers=auth());assert r.status_code==200,r.text
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PublicKey
    from server.domain import canonical
    envelope=r.json();Ed25519PublicKey.from_public_bytes(base64.b64decode(envelope['publicKey'])).verify(base64.b64decode(envelope['signature']),canonical(envelope['manifest']).encode())
    from docx import Document
    from openpyxl import load_workbook
    d=c.get(url+'docx',headers=auth());assert d.status_code==200;assert any('مهندس' in p.text for p in Document(io.BytesIO(d.content)).paragraphs)
    x=c.get(url+'xlsx',headers=auth());assert x.status_code==200;w=load_workbook(io.BytesIO(x.content));assert set(w.sheetnames)=={'Position','RACI','Skills'};assert w['RACI']['C2'].value=='Manager'
    pdf=c.get(url+'pdf',headers=auth());assert pdf.status_code==200;assert pdf.content.startswith(b'%PDF')

def test_login_revocation_and_failed_attempt_limit(env):
    app,c,auth,create=env
    for _ in range(5):assert c.post('/api/v1/auth/login',json={'email':'unknown@example.test','password':'bad'}).status_code==401
    assert c.post('/api/v1/auth/login',json={'email':'unknown@example.test','password':'bad'}).status_code==429
    assert c.post('/api/v1/auth/logout-all',headers=auth()).status_code==200
    assert c.get('/api/v1/me',headers=auth()).status_code==401

def test_concurrent_approvals_advance_one_stage_only(env):
    import os
    if not os.getenv('MIYAR_TEST_POSTGRES_URL'):pytest.skip('Row-lock concurrency gate runs against PostgreSQL in CI')
    from concurrent.futures import ThreadPoolExecutor
    app,c,auth,create=env;p=create();submit(c,auth,p)
    def approve():return decide(c,auth,p,'od_specialist',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01'}).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:statuses=list(pool.map(lambda _:approve(),range(2)))
    assert statuses.count(200)==1;assert all(s in [200,403,409] for s in statuses)
    assert c.get('/api/v1/positions/'+p['id'],headers=auth()).json()['approvalStage']==1

def test_organization_rules_are_enforced_at_approval_and_snapshotted(env):
    app,c,auth,create=env
    policy={'similarityThreshold':.8,'minimumMargin':.02,'rules':{'251204':{'minimumEducationLevel':6,'licenseRequired':True}}}
    assert c.post('/api/v1/settings/decision-policy',headers=auth('admin'),json={'policy':policy}).status_code==200
    p=create(educationLevel='6');submit(c,auth,p)
    # A later policy change cannot erase the requirements already submitted for review.
    assert c.post('/api/v1/settings/decision-policy',headers=auth('admin'),json={'policy':{**policy,'rules':{}}}).status_code==200
    assert decide(c,auth,p,'od_specialist',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01'}).status_code==422
    assert decide(c,auth,p,'od_specialist',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01','licenseVerifiedByOD':True}).status_code==200

def test_free_text_field_seniority_roundtrip_and_exports(env):
    app,c,auth,create=env
    p=create(field='Health informatics / AI governance',seniority='Principal - 12+ years')
    value=c.get('/api/v1/positions/'+p['id']+'/export/json',headers=auth()).json()
    assert value['content']['field']=='Health informatics / AI governance'
    assert value['content']['seniority']=='Principal - 12+ years'
    from docx import Document
    r=c.get('/api/v1/positions/'+p['id']+'/export/docx',headers=auth())
    assert any('Principal - 12+ years' in x.text for x in Document(io.BytesIO(r.content)).paragraphs)
    assert c.get('/api/v1/capabilities').status_code==401
    assert c.get('/api/v1/capabilities',headers=auth()).json()['approvals'] is True

def test_approved_evaluation_is_locked_after_rewards_stage(env):
    app,c,auth,create=env;p=create();submit(c,auth,p)
    decide(c,auth,p,'od_specialist',{'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Department reviewer for test','businessReviewDate':'2026-09-01'})
    url='/api/v1/positions/'+p['id']+'/evaluation';body={'revision':1,'answers':dict.fromkeys(['knowledge','complexity','impact'],'2'),'evidence':dict.fromkeys(['knowledge','complexity','impact'],'Reviewed scope')}
    assert c.post(url,headers=auth('total_rewards'),json=body).status_code==200
    assert decide(c,auth,p,'total_rewards',{'payFrameworkReviewed':True}).status_code==200
    body['answers']['impact']='4'
    assert c.post(url,headers=auth('total_rewards'),json=body).status_code==409

def test_invalid_matrix_scope_and_provisional_code_return_validation_errors(env):
    app,c,auth,create=env;p=create()
    for change in [{'raci':[1]},{'skillRequirements':[{'name':['bad']}]},{'directReports':1.5},{'provisional':True}]:
        assert c.patch('/api/v1/positions/'+p['id'],headers=auth(),json={'revision':1,'reason':'Test validation','content':{**p['content'],**change}}).status_code==422
