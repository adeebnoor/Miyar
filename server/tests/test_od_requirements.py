"""Synthetic OD workflow cases. No customer source documents or proprietary tables."""
import copy,io
from datetime import date
import pytest
from docx import Document
from openpyxl import load_workbook
from server.domain import validate_framework,grade
from .conftest import submit,decide

OD={'scopeReviewed':True,'mappingReviewed':True,'businessValidated':True,'roleNotPerson':True,'businessReviewer':'Synthetic department reviewer','businessReviewDate':'2026-09-01'}
FRAMEWORK={'id':'synthetic-external','name':'Synthetic authorized external report','method':'external-korn-ferry-record','licenseReference':'TEST-AUTHORIZATION-NOT-A-REAL-LICENSE','approvedBy':'Synthetic authority','approvedOn':'2026-09-01','authorizedUseConfirmed':True}
EVIDENCE={'reportReference':'SYNTHETIC-REPORT-ONLY','assessor':'Synthetic specialist','evaluationDate':'2026-09-01','rationale':'Synthetic report justification based on reviewed position scope','knowledge':'Synthetic knowledge evidence','problemSolving':'Synthetic problem-solving evidence','accountability':'Synthetic accountability evidence','reportConfirmed':True}

def test_external_record_rejects_unverified_or_incomplete_report_inputs():
    for change in [{'authorizedUseConfirmed':False},{'licenseReference':''},{'approvedOn':'2999-01-01'}]:
        with pytest.raises(ValueError):validate_framework({**FRAMEWORK,**change})
    for change in [{'reportConfirmed':False},{'rationale':''},{'evaluationDate':'2026-02-30'}]:
        with pytest.raises(ValueError):grade(FRAMEWORK,{'score':412,'band':'Test grade'},{**EVIDENCE,**change})
    for score in [True,float('nan'),-1,'412']:
        with pytest.raises(ValueError):grade(FRAMEWORK,{'score':score,'band':'Test grade'},EVIDENCE)
    result=grade(FRAMEWORK,{'score':412,'band':'Test grade'},EVIDENCE)
    assert result['computedBy']=='external-specialist-report' and result['points']==412
    assert result['externalReport']['rationale']==EVIDENCE['rationale']

def test_department_consultation_and_regulatory_evidence_are_required_before_od_approval(env):
    app,c,auth,create=env
    p=create(certifications='Synthetic certificate',saudization='Synthetic policy requirement',saudizationSource='https://example.test/policy',saudizationDate='2026-09-01',license='Synthetic license requirement',licenseSource='https://example.test/license',licenseDate='2026-09-01',educationFieldCode='071501')
    submit(c,auth,p)
    for change in [{'businessValidated':False},{'roleNotPerson':False},{'businessReviewDate':'2999-01-01'},{'businessReviewer':''},{'regulatoryReviewed':False}]:
        assert decide(c,auth,p,'od_specialist',{**OD,'regulatoryReviewed':True,**change}).status_code==422
    assert decide(c,auth,p,'od_specialist',{**OD,'regulatoryReviewed':True}).status_code==200
    for invalid in [{'licenseSource':'javascript:alert(1)'},{'licenseDate':'2026-02-30'},{'saudizationDate':'2999-01-01'}]:
        r=c.post('/api/v1/positions',headers=auth(),json={'departmentId':'org-a-one','content':{**p['content'],**invalid},'reason':'Invalid policy fixture'});assert r.status_code==422
    q=create(saudization='Recorded without evidence');submit(c,auth,q)
    assert decide(c,auth,q,'od_specialist',{**OD,'regulatoryReviewed':True}).status_code==422

def test_authorized_specialist_report_completes_workflow_and_exports_complete_evidence(env):
    app,c,auth,create=env
    assert c.post('/api/v1/settings/framework',headers=auth('finance'),json={'framework':FRAMEWORK,'reason':'Synthetic authorization'}).status_code==403
    assert c.post('/api/v1/settings/framework',headers=auth('admin'),json={'framework':FRAMEWORK,'reason':'Synthetic authorization'}).status_code==200
    p=create(certifications='Synthetic certificate requirement',educationFieldCode='071501',directReports=0,effectiveDate='2027-01-01');url='/api/v1/positions/'+p['id']
    body={'revision':1,'answers':{'score':412,'band':'Synthetic grade'},'evidence':EVIDENCE}
    assert c.post(url+'/evaluation',headers=auth('total_rewards'),json=body).status_code==409
    submit(c,auth,p);assert decide(c,auth,p,'od_specialist',{**OD,'regulatoryReviewed':True}).status_code==200
    assert c.post(url+'/evaluation',headers=auth('finance'),json=body).status_code==403
    assert c.post(url+'/evaluation',headers=auth('total_rewards'),json=body).status_code==200
    assert decide(c,auth,p,'total_rewards',{'payFrameworkReviewed':True}).status_code==200
    assert decide(c,auth,p,'finance',{'vacancyConfirmed':True,'budgetConfirmed':True,'approvedAnnualBudget':240000,'approvedHeadcount':1}).status_code==200
    assert decide(c,auth,p,'chro',{}).status_code==200
    assert c.post(url+'/evaluation',headers=auth('total_rewards'),json=body).status_code==409
    result=c.get(url+'/export/json?lang=en',headers=auth()).json()
    assert result['approved'] and result['evaluation']['result']['points']==412
    assert result['approvals'][0]['evidence']['businessReviewer']==OD['businessReviewer']
    doc=c.get(url+'/export/docx?lang=en',headers=auth());assert doc.status_code==200
    text='\n'.join(p.text for p in Document(io.BytesIO(doc.content)).paragraphs)
    for value in ['071501','2027-01-01','Synthetic certificate requirement',EVIDENCE['rationale'],OD['businessReviewer'],'not calculated']:assert value in text
    excel=c.get(url+'/export/xlsx?lang=en',headers=auth());assert excel.status_code==200
    sheet=load_workbook(io.BytesIO(excel.content))['Position'];values=[str(cell.value) for row in sheet for cell in row]
    assert EVIDENCE['rationale'] in values and '071501' in values and OD['businessReviewer'] in values
    pdf=c.get(url+'/export/pdf?lang=en',headers=auth());assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
