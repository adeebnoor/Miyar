import csv,io,json,math,os,re,threading,zipfile,hashlib
from collections import Counter,defaultdict
from pathlib import Path
import numpy as np
from .domain import normalized,digest

ROOT=Path(__file__).resolve().parent.parent
class Catalog:
    def __init__(self,directory=None,occupation_payload=None):
        path=Path(directory or ROOT/'dist/classifications')
        self.occupations=occupation_payload or json.loads((path/'ssco-2019.json').read_text());self.education=json.loads((path/'education-2020.json').read_text())
        self.nodes={n['code']:n for n in self.occupations['nodes']};self.roles={k:v for k,v in self.nodes.items() if v['level']=='occupation'}
        self.education_fields={n['code']:n for n in self.education['fields']}
        self.flagged={n['code'] for n in self.occupations['validation']['missingParents']}
        self.skills=json.loads((path/'skills.json').read_text()) if (path/'skills.json').exists() else []
        self.profiles={p['code']:p for p in json.loads((path/'role-profiles.json').read_text())} if (path/'role-profiles.json').exists() else {}
        if occupation_payload:self.profiles={}
        for k,p in self.profiles.items():
            if k in self.roles:self.roles[k]['titleEn']=p['titleEn']
        self.role_skills={k:p['skills'] for k,p in self.profiles.items()};self.model=None;self.matrix=None;self.model_lock=threading.Lock()
    def search(self,text='',parent=None,limit=50):
        q=normalized(text);parts=q.split();rows=self.nodes.values() if parent is not None else self.roles.values()
        found=[r for r in rows if (parent is None or r.get('parent')==parent) and all(t in normalized(r['titleAr']+' '+r['code']+' '+r.get('titleEn','')+' '+' '.join(r.get('aliases',[]))) for t in parts)]
        found.sort(key=lambda r:(normalized(r['titleAr'])!=q and r['code']!=q,r['code']))
        return {'total':len(found),'items':[{**r,'sourceWarning':r['code'] in self.flagged} for r in found[:limit]],'release':self.occupations['id']}
    def extract_skills(self,text):
        normalized_text=' '+normalized(text)+' ';found=[]
        for skill in self.skills:
            terms=[skill.get('labelAr',''),skill.get('labelEn',''),*skill.get('aliases',[])]
            evidence=[term for term in terms if term and any(prefix+normalized(term)+' ' in normalized_text for prefix in [' ',' و',' ف'])]
            if evidence:found.append({**skill,'evidence':list(dict.fromkeys(evidence)),'method':'dictionary-extraction'})
        return found
    def load_semantic(self):
        name=os.getenv('MIYAR_EMBEDDING_MODEL','sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2')
        if os.getenv('MIYAR_ENABLE_EMBEDDINGS')!='true':raise RuntimeError('Semantic model is disabled; configure MIYAR_ENABLE_EMBEDDINGS=true and provision model resources')
        with self.model_lock:
            if self.model is not None:return
            from fastembed import TextEmbedding
            from huggingface_hub import snapshot_download
            revision='faf4aa4225822f3bc6376869cb1164e8e3feedd0'
            folder=snapshot_download('qdrant/paraphrase-multilingual-MiniLM-L12-v2-onnx-Q',revision=revision,cache_dir=os.getenv('MIYAR_MODEL_CACHE',str(ROOT/'.model-cache')))
            model=TextEmbedding(name,specific_model_path=folder,threads=2)
            artifacts={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in Path(folder).iterdir() if p.is_file()}
            self.model_fingerprint=digest({'artifacts':artifacts,'fastembed':'0.8.0','pooling':'mean'})
            roles=list(self.roles.values());documents=[r['titleAr']+' '+self.nodes.get(r.get('parent'),{}).get('titleAr','')+' '+r.get('titleEn','') for r in roles]
            cache=Path(os.getenv('MIYAR_MODEL_CACHE',str(ROOT/'.model-cache')));cache.mkdir(parents=True,exist_ok=True)
            filename=cache/('occupation-vectors-'+digest({'release':self.occupations['sha256'],'model':name,'fingerprint':self.model_fingerprint,'text':documents})[:24]+'.npy')
            matrix=np.load(filename,allow_pickle=False) if filename.exists() else np.asarray(list(model.embed(documents,batch_size=32)))
            matrix=matrix/np.maximum(np.linalg.norm(matrix,axis=1,keepdims=True),1e-12)
            if not filename.exists():np.save(filename,matrix,allow_pickle=False)
            self.model=model;self.matrix=matrix;self.vector_roles=roles;self.model_name=name
    def semantic(self,text,candidate_codes=None):
        self.load_semantic();vector=np.asarray(next(iter(self.model.embed([text]))));vector/=max(float(np.linalg.norm(vector)),1e-12)
        cosine=self.matrix@vector;order=np.argsort(-cosine)
        if candidate_codes is not None:order=[i for i in order if self.vector_roles[int(i)]['code'] in candidate_codes]
        order=order[:3];skills=self.extract_skills(text);required={s['id'] for s in skills};items=[]
        for index in order:
            role=self.vector_roles[int(index)];known=set(self.role_skills.get(role['code'],[]));common=sorted(required&known)
            items.append({'code':role['code'],'titleAr':role['titleAr'],'sourcePage':role['sourcePage'],'cosineSimilarity':round(float(cosine[index]),5),'semanticDistance':round(1-float(cosine[index]),5),'skillOverlapPercent':round(100*len(common)/len(required),1) if required and known else None,'matchedSkillIds':common,'sourceWarning':role['code'] in self.flagged,'skillDenominator':len(required),'profile':self.profiles.get(role['code']),'taskOverlapPercent':None})
        return {'method':'multilingual-embedding-cosine','model':self.model_name,'modelFingerprint':self.model_fingerprint,'referenceText':'Occupation title and unit title; the supplied SSCO PDF does not contain task/skill profiles','release':self.occupations['id'],'candidates':items,'extractedSkills':skills,'notice':'Cosine similarity is not a probability, approval, or a 65/35 role composition. Skill overlap is the share of extracted request skills in the selected proposed O*NET crosswalk. These two authored crosswalks require OD review; absent profiles and task overlap remain null.'}

def read_rows(raw,filename):
    if len(raw)>10_000_000:raise ValueError('Upload limit: 10 MB')
    if filename.lower().endswith('.csv'):
        text=raw.decode('utf-8-sig');reader=csv.DictReader(io.StringIO(text),strict=True);rows=[]
        if not reader.fieldnames or len(reader.fieldnames)!=len(set(reader.fieldnames)):raise ValueError('Missing or duplicate column names')
        for i,row in enumerate(reader):
            if i>=10000:raise ValueError('Row limit: 10,000')
            if None in row:raise ValueError('Row has more values than the header')
            rows.append(row)
    elif filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist())>100_000_000:raise ValueError('Expanded workbook limit exceeded')
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False);sheet=workbook.active;values=sheet.iter_rows(values_only=True);headers=[str(x or '').strip() for x in next(values,[])];rows=[]
        if len(headers)!=len(set(headers)):raise ValueError('Duplicate column names')
        for i,row in enumerate(values):
            if i>=10000:raise ValueError('Row limit: 10,000')
            rows.append(dict(zip(headers,row)))
        workbook.close()
    else:raise ValueError('Use UTF-8 CSV or XLSX')
    aliases={'المسمى':'title','المسمى الوظيفي':'title','الإدارة':'department','الرمز المهني':'occupationCode','المسؤوليات':'responsibilities','المرؤوسون':'directReports','الميزانية':'budgetAmount','الصلاحيات':'authority','العدد':'headcount'}
    allowed={'title','department','occupationCode','responsibilities','directReports','budgetAmount','authority','headcount','internalCode'}
    headers=reader.fieldnames if filename.lower().endswith('.csv') else headers
    canonical_headers=[aliases.get(str(k).strip(),str(k).strip()) for k in headers]
    if len(canonical_headers)!=len(set(canonical_headers)):raise ValueError('Duplicate column names after language normalization')
    clean=[]
    for row in rows:
        item={aliases.get(str(k).strip(),str(k).strip()):str('' if v is None else v).strip() for k,v in row.items() if k is not None}
        if not any(item.values()):continue
        if not item.get('title'):raise ValueError('Every non-empty row needs a title / المسمى')
        if any(len(v)>10000 for v in item.values()):raise ValueError('Cell length exceeds 10,000 characters')
        clean.append({k:v for k,v in item.items() if k in allowed})
    if not clean:raise ValueError('No position rows found')
    return clean

def bulk_diagnosis(rows,catalog):
    output=[];groups=defaultdict(list);matched=valid=assessable=inflated=0
    for index,row in enumerate(rows,2):
        code=normalized(row.get('occupationCode','')).replace(' ','');record=catalog.roles.get(code);title=normalized(row['title']);exact=record is not None and title==normalized(record['titleAr'])
        source_ok=record is not None and code not in catalog.flagged
        valid+=int(source_ok);matched+=int(exact and source_ok);flags=[]
        if not code:flags.append('missing_occupation_code')
        elif not record:flags.append('code_not_in_selected_release')
        elif code in catalog.flagged:flags.append('source_hierarchy_issue')
        elif not exact:flags.append('title_code_pair_needs_review')
        scope=all(row.get(k,'')!='' for k in ['directReports','budgetAmount','authority'])
        if scope:
            try:
                numbers=[float(row['directReports']),float(row['budgetAmount'])]
                if any(not math.isfinite(n) or n<0 for n in numbers):raise ValueError('Invalid scope')
                assessable+=1
                senior=bool(re.search(r'\b(director|head|chief|مدير|رئيس)\b',title));limited=float(row['directReports'])==0 and float(row['budgetAmount'])==0 and any(t in normalized(row['authority']) for t in ['recommend','يقترح','توصيات','يوصي'])
                if senior and limited:flags.append('title_scope_review');inflated+=int(senior and limited)
            except (ValueError,OverflowError):flags.append('invalid_scope_numbers')
        groups[(title,normalized(row.get('department','')))].append(index)
        output.append({'row':index,'title':row['title'],'department':row.get('department',''),'occupationCode':code or None,'sourceTitle':record['titleAr'] if record else None,'sourcePage':record['sourcePage'] if record else None,'titleCodeAligned':exact and source_ok,'flags':flags})
    duplicates=[{'title':key[0],'department':key[1],'rows':values} for key,values in groups.items() if len(values)>1]
    return {'release':catalog.occupations['id'],'totalRows':len(rows),'codesFoundWithoutSourceIssue':valid,'exactTitleCodePairs':matched,'titleCodeAlignmentPercent':round(100*matched/len(rows),1),'duplicateGroups':duplicates,'scopeAssessableRows':assessable,'titleScopeReviewRows':inflated,'rows':output,'notice':'Alignment measures exact title/code pairs in the selected edition, not legal compliance or task fit. Repeated titles may represent legitimate additional headcount; title-scope flags are advisory, not proof of inflation. Employee identity columns are discarded.'}
