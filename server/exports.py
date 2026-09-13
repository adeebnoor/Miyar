"""Portable exports. Internal cryptographic approval receipts are not qualified PDF signatures."""
import base64,html,io,os
from .domain import canonical,digest

LABELS={'field':('المجال الوظيفي','Field'),'seniority':('المستوى المطلوب','Seniority'),'headcount':('العدد المطلوب','Requested headcount'),'annualCost':('التكلفة السنوية الإجمالية - ر.س','Total annual cost - SAR'),'manager':('المدير المباشر','Reporting manager'),'title':('المسمى الوظيفي','Job title'),'businessNeed':('المبرر التجاري','Business need'),'alternatives':('البدائل المدروسة','Alternatives considered'),'successMeasures':('مؤشرات النجاح','Success measures'),'purpose':('هدف المنصب','Position purpose'),'responsibilities':('المسؤوليات','Responsibilities'),'team':('الفريق ونطاق الإشراف','Team and supervision'),'budget':('نطاق الميزانية','Budget scope'),'authority':('الصلاحيات','Decision authority'),'impact':('نطاق الأثر','Impact'),'stakeholders':('أصحاب العلاقة','Stakeholders'),'qualifications':('المؤهلات','Qualifications'),'experience':('الخبرة','Experience'),'skills':('المهارات','Skills'),'behaviors':('الجدارات السلوكية','Behaviors'),'occupationCode':('رمز المهنة','Occupation code'),'occupationRelease':('إصدار التصنيف','Classification edition'),'educationLevel':('المستوى التعليمي','Education level'),'educationFieldCode':('رمز التخصص التعليمي','Education field code'),'constraints':('المحددات','Constraints')}
LABELS.update({'department': ('الإدارة', 'Department'), 'directReports': ('عدد المرؤوسين', 'Direct reports'), 'certifications': ('الشهادات المطلوبة', 'Required certifications'), 'saudization': ('متطلب التوطين المدخل', 'Recorded Saudization requirement'), 'saudizationSource': ('مصدر متطلب التوطين', 'Saudization source URL'), 'saudizationDate': ('تاريخ التحقق من التوطين', 'Saudization check date'), 'license': ('متطلب الترخيص المهني المدخل', 'Recorded professional licensing requirement'), 'licenseSource': ('مصدر الترخيص المهني', 'Licensing source URL'), 'licenseDate': ('تاريخ التحقق من الترخيص', 'Licensing check date'), 'effectiveDate': ('تاريخ السريان المقترح', 'Proposed effective date'), 'sourceDecisionId': ('مرجع القرار المستورد', 'Imported decision reference'), 'sourceDecisionInput': ('مدخلات القرار الأصلي', 'Original decision inputs'), 'importNotes': ('ملاحظات الاستيراد', 'Import review notes')})

def signing_key():
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    raw=os.getenv('MIYAR_SIGNING_KEY','')
    if not raw:raise ValueError('Configure a persistent base64 Ed25519 signing key to issue approval receipts')
    return Ed25519PrivateKey.from_private_bytes(base64.b64decode(raw,validate=True))
def public_key():
    from cryptography.hazmat.primitives.serialization import Encoding,PublicFormat
    data=signing_key().public_key().public_bytes(Encoding.Raw,PublicFormat.Raw)
    return {'algorithm':'Ed25519','publicKey':base64.b64encode(data).decode(),'keyId':digest({'publicKey':base64.b64encode(data).decode()})[:16]}
def receipt(card):
    if not card['approved']:raise ValueError('Approval receipts are only issued for approved revisions')
    key=public_key();manifest={'positionId':card['id'],'internalCode':card['internalCode'],'revision':card['revision'],'contentDigest':digest(card['content']),'evaluationDigest':digest(card.get('evaluation')),'approvals':card['approvals'],'keyId':key['keyId']}
    return {'manifest':manifest,'signature':base64.b64encode(signing_key().sign(canonical(manifest).encode())).decode(),**key,'notice':'Internal organization approval receipt. This is not a qualified electronic signature or a PAdES-signed PDF.'}
def safe_cell(value):
    # Exports contain untrusted user input; Excel must treat it as literal text.
    return str(value if value is not None else '')
def matrices(card):
    return [('RACI',['Responsibility / المسؤولية','R — Responsible','A — Accountable','C — Consulted','I — Informed'],[[r.get(k,'') for k in ['responsibility','R','A','C','I']] for r in card['content'].get('raci',[])]),('Skills',['Skill / المهارة','Type / النوع','Required level / المستوى','Evidence / الدليل'],[[r.get(k,'') for k in ['name','type','level','evidence']] for r in card['content'].get('skillRequirements',[])])]
def evaluation_lines(card):
    ar=card['lang']=='ar';rows=[]
    for approval in card.get('approvals',[]):
        if approval.get('role')=='od_specialist' and approval.get('decision')=='approve':
            evidence=approval.get('evidence',{})
            rows += [('مراجع الإدارة' if ar else 'Department reviewer',evidence.get('businessReviewer','')),('تاريخ مراجعة الإدارة' if ar else 'Department review date',evidence.get('businessReviewDate','')),('التحقق من وصف المنصب' if ar else 'Role-based description confirmed',str(evidence.get('roleNotPerson',False)))]
    if not card.get('evaluation'):return rows
    e=card['evaluation']['result']
    rows += [('طريقة التقييم' if ar else 'Evaluation method',e.get('method','custom'))]
    if e.get('externalReport'):
        labels={'reportReference':('مرجع تقرير المختص','Specialist report reference'),'assessor':('المقيّم الخارجي','External assessor'),'evaluationDate':('تاريخ التقييم','Evaluation date'),'rationale':('مبررات التقييم','Evaluation rationale'),'knowledge':('دليل المعرفة','Knowledge evidence'),'problemSolving':('دليل حل المشكلات','Problem-solving evidence'),'accountability':('دليل المساءلة','Accountability evidence')}
        rows += [(labels[k][0 if ar else 1],v) for k,v in e['externalReport'].items() if k in labels]
        rows.append(('أساس النتيجة' if ar else 'Result basis','نتيجة أدخلها مختص الجهة من تقرير خارجي؛ لم يحسبها معيار ولم يتحقق من ترخيصها مستقلًا.' if ar else 'Recorded by the organization specialist from an external report; not calculated or independently license-verified by Miyar.'))
    else:
        rows += [(('مبرر العامل ' if ar else 'Factor rationale ')+str(x['factor']),str(x.get('evidence',''))) for x in e.get('breakdown',[])]
    return rows

def xlsx(card):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    w=Workbook();w.remove(w.active)
    metadata=[['Internal code',card['internalCode']],['Revision',card['revision']],['State','Approved' if card['approved'] else 'Draft — not approved'],['Source',card['content'].get('occupationRelease','')]]
    if card.get('evaluation'):
        e=card['evaluation']['result'];metadata.extend([['Evaluation points',e['points']],['Evaluation band',e['band']['id']],['Framework',str(e['frameworkId'])+' / '+str(e['frameworkVersion'])]])
    metadata.extend(evaluation_lines(card))
    for name,headers,rows in [('Position',['Field / الحقل','Value / القيمة'],metadata+[[v[1]+' / '+v[0],card['content'].get(k,'')] for k,v in LABELS.items()]),*matrices(card)]:
        s=w.create_sheet(name);s.sheet_view.rightToLeft=card['lang']=='ar';s.append(headers)
        for row in rows:
            s.append([safe_cell(v) for v in row])
            for cell in s[s.max_row]:cell.data_type='s'
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for row in s:
            for c in row:
                c.font=Font(name='DejaVu Sans',size=11,color='FFFFFF' if c.row==1 else '172E35',bold=c.row==1)
                c.fill=PatternFill('solid',fgColor='183E48' if c.row==1 else ('F1F6F5' if c.row%2 else 'FFFFFF'))
                c.alignment=Alignment(vertical='top',wrap_text=True)
                c.border=Border(bottom=Side(style='hair',color='D9D9D9'))
            s.row_dimensions[row[0].row].height=max(30,min(180,max(len(str(c.value or '')) for c in row)//35*16+24))
        for i in range(1,len(headers)+1):s.column_dimensions[get_column_letter(i)].width=58 if i==1 or name=='Position' else 28
        s.print_options.horizontalCentered=True;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    out=io.BytesIO();w.save(out);return out.getvalue()
def docx(card,brand):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();section=d.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11);section.top_margin=section.bottom_margin=Inches(.7)
    ar=card['lang']=='ar';normal=d.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(8)
    for name in ['Title','Heading 1','Heading 2']:d.styles[name].font.color.rgb=RGBColor(0,0,0)
    def para(text,style=None):
        p=d.add_paragraph(str(text),style)
        if ar:
            bidi=OxmlElement('w:bidi');p._p.get_or_add_pPr().append(bidi)
        return p
    section.header.paragraphs[0].text=brand.get('nameAr' if ar else 'nameEn','Miyar')
    para('بطاقة الوصف الوظيفي' if ar else 'Job description','Title');para(card['content']['title'],'Heading 1')
    para(f"{card['internalCode']} • v{card['revision']} • "+(('معتمد' if ar else 'Approved') if card['approved'] else ('مسودة غير معتمدة' if ar else 'Unapproved draft')))
    for k,label in LABELS.items():
        if k!='title' and card['content'].get(k) is not None and card['content'].get(k)!='':para(label[0 if ar else 1],'Heading 2');para(card['content'][k])
    if card.get('evaluation'):
        e=card['evaluation']['result'];para('التقييم الوظيفي' if ar else 'Job evaluation','Heading 2');para(str(e['points'])+' / '+e['band']['id']+' | '+str(e['frameworkId'])+' v'+str(e['frameworkVersion']))
    for title,value in evaluation_lines(card):para(title,'Heading 2');para(value)
    for name,headers,rows in matrices(card):
        para(name,'Heading 2');table=d.add_table(rows=1,cols=len(headers));table.style='Light Shading Accent 1'
        for i,h in enumerate(headers):table.rows[0].cells[i].text=h
        for row in rows:
            for c,v in zip(table.add_row().cells,row):c.text=str(v)
        if not rows:para('لم تُدخل بيانات هذه المصفوفة' if ar else 'No matrix entries supplied')
        for row in table.rows:
            for cell in row.cells:
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                cell._tc.get_or_add_tcPr().append(borders)
    para('سجل الاعتماد' if ar else 'Approval record','Heading 2')
    for a in card['approvals']:para(a['actorName']+' • '+a['role']+' • '+a['createdAt'])
    section.footer.paragraphs[0].text=brand.get('footer','')+' | '+card['internalCode']
    out=io.BytesIO();d.save(out);return out.getvalue()
def pdf(card,brand):
    from weasyprint import HTML
    ar=card['lang']=='ar';esc=lambda v:html.escape(str(v));color=brand.get('color','#146954')
    fields=''.join('<section><h2>'+esc(label[0 if ar else 1])+'</h2><p>'+esc(card['content'][k]).replace('\n','<br>')+'</p></section>' for k,label in LABELS.items() if k!='title' and card['content'].get(k) is not None and card['content'].get(k)!='')
    if card.get('evaluation'):
        e=card['evaluation']['result'];fields+='<section><h2>'+('التقييم الوظيفي' if ar else 'Job evaluation')+'</h2><p>'+esc(str(e['points'])+' / '+e['band']['id']+' | '+str(e['frameworkId'])+' v'+str(e['frameworkVersion']))+'</p></section>'
    fields+=''.join('<section><h2>'+esc(title)+'</h2><p>'+esc(value).replace('\n','<br>')+'</p></section>' for title,value in evaluation_lines(card))
    for name,headers,rows in matrices(card):
        if rows:fields+='<section><h2>'+esc(name)+'</h2><table><thead><tr>'+''.join('<th>'+esc(h)+'</th>' for h in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(v)+'</td>' for v in row)+'</tr>' for row in rows)+'</tbody></table></section>'
    approvals=''.join('<tr><td>'+esc(a['actorName'])+'</td><td>'+esc(a['role'])+'</td><td dir="ltr">'+esc(a['createdAt'][:19])+'</td></tr>' for a in card['approvals'])
    seal=card.get('receipt');seal_text=('بصمة سجل الاعتماد' if ar else 'Approval manifest digest')+': '+digest(seal['manifest']) if seal else ('لا يوجد ختم اعتماد لهذا الملف' if ar else 'No approval seal for this file')
    markup=f'''<!doctype html><html lang="{card['lang']}" dir="{'rtl' if ar else 'ltr'}"><meta charset="utf-8"><style>@page{{size:A4;margin:22mm 20mm;@bottom-center{{content:counter(page);font:9pt sans-serif}}}}body{{font:11pt 'DejaVu Sans',sans-serif;color:#172e35;line-height:1.7}}header{{color:{color};font-size:13pt;border-bottom:2pt solid {color};padding-bottom:12pt}}h1{{font-size:23pt;line-height:1.4}}h2{{font-size:12pt;margin:12pt 0 4pt;break-after:avoid}}p{{margin:3pt 0;white-space:normal;overflow-wrap:anywhere}}section{{orphans:3;widows:3}}table{{width:100%;border-collapse:collapse;font-size:9pt}}th,td{{padding:6pt;border:1px solid #d9d9d9}}.meta{{font-size:10pt;color:#566e72}}footer{{break-inside:avoid;margin-top:12pt;font-size:8pt;overflow-wrap:anywhere}}</style><header>{esc(brand.get('nameAr' if ar else 'nameEn','Miyar'))}</header><h1>{esc(card['content']['title'])}</h1><p class="meta" dir="ltr">{esc(card['internalCode'])} · v{card['revision']} · {'APPROVED' if card['approved'] else 'DRAFT — NOT APPROVED'}</p>{fields}<h2>{'سجل الاعتماد' if ar else 'Approval record'}</h2><table>{approvals}</table><footer><p>{esc(seal_text)}</p><p>{'سجل اعتماد داخلي موثق في النظام؛ ليس توقيع PDF مؤهلاً قانونيًا.' if ar else 'Internal system approval record; not a qualified PDF signature.'}</p><p>{esc(brand.get('footer',''))}</p></footer></html>'''
    def deny_fetch(url,*args,**kwargs):raise ValueError('External resources are disabled for document exports')
    return HTML(string=markup,url_fetcher=deny_fetch).write_pdf()
